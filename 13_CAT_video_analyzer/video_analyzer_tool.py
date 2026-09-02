import json
import os
import re
import time
import zipfile
from io import BytesIO
from datetime import datetime

import httpx
from cat.log import log
from cat.mad_hatter.decorators import hook, tool
from docx.shared import Inches

from .video_analyzer_client import delete_video_job
from .result_cleaner import clean_result_string
from .session_store import create_session, delete_session, get_session
from .word_formatters import load_template_document, parse_markdown_to_word


UPLOAD_PAGE_URL = "http://127.0.0.1:1865/video-analyzer/upload"
WAIT_TIMEOUT_SECONDS = 1800
POLL_INTERVAL_SECONDS = 2


def _load_tools_status() -> dict:
    try:
        with open("cat/static/tools_status.json", "r", encoding="utf-8") as f:
            return json.load(f) or {}
    except Exception:
        return {}


def _is_tool_enabled_for_user(cat, tool_key: str) -> bool:
    ts = _load_tools_status()
    uid = str(getattr(cat, "user_id", "") or "")
    return bool(
        ts.get("tools", {})
        .get(tool_key, {})
        .get("user_id_tool_status", {})
        .get(uid, False)
    )


def _get_backend_base_url(cat) -> str:
    settings = cat.mad_hatter.get_plugin().load_settings()
    return settings.get("backend_base_url", "").rstrip("/")


def _delete_backend_job_if_possible(cat, job_id: str | None):
    if not job_id:
        return

    backend_base_url = _get_backend_base_url(cat)
    if not backend_base_url:
        return

    try:
        delete_video_job(job_id=job_id, backend_base_url=backend_base_url)
    except Exception as e:
        log.warning(f"[Video Analyzer Tool] Cleanup job fallita per {job_id}: {e}")


def _extract_test_job_id(result_text: str) -> str | None:
    try:
        payload = json.loads(result_text)
        return payload.get("job_id") or payload.get("result", {}).get("job_id")
    except json.JSONDecodeError:
        pass

    match = re.search(r'"job_id"\s*:\s*"([^"]+)"', result_text)
    if match:
        return match.group(1)

    return None


def _normalize_backend_url(backend_base_url: str, raw_url: str) -> str:
    if not raw_url:
        return ""
    if raw_url.startswith("http://") or raw_url.startswith("https://"):
        return raw_url
    return f"{backend_base_url}{raw_url if raw_url.startswith('/') else '/' + raw_url}"


def _collect_frame_urls(node, backend_base_url: str, collected: list[str]):
    if isinstance(node, str):
        lower_node = node.lower()
        if "/frames/" in lower_node or lower_node.endswith((".png", ".jpg", ".jpeg", ".webp")):
            collected.append(_normalize_backend_url(backend_base_url, node))
        return

    if isinstance(node, list):
        for item in node:
            _collect_frame_urls(item, backend_base_url, collected)
        return

    if not isinstance(node, dict):
        return

    for key in ("url", "image_url", "frame_url", "download_url", "path", "file", "filename"):
        value = node.get(key)
        if isinstance(value, str):
            lower_value = value.lower()
            if "/frames/" in lower_value or lower_value.endswith((".png", ".jpg", ".jpeg", ".webp")):
                collected.append(_normalize_backend_url(backend_base_url, value))

    for value in node.values():
        _collect_frame_urls(value, backend_base_url, collected)


def _frame_sort_key(frame_url: str):
    match = re.search(r"frame[_-]?(\d+)", frame_url, re.IGNORECASE)
    if match:
        return int(match.group(1))
    return 10**9


def _fetch_backend_frame_urls(job_id: str, backend_base_url: str) -> list[str]:
    if not job_id or not backend_base_url:
        return []

    frames_url = f"{backend_base_url}/api/jobs/{job_id}/frames"
    response = httpx.get(frames_url, timeout=60.0)
    response.raise_for_status()

    payload = response.json()
    collected: list[str] = []
    _collect_frame_urls(payload, backend_base_url, collected)
    unique_urls = sorted(set(filter(None, collected)), key=_frame_sort_key)
    return unique_urls


def _download_frame_images_to_static(frame_urls: list[str], job_id: str) -> dict[str, str]:
    frame_images: dict[str, str] = {}
    frames_dir = os.path.join("cat", "static", "video_prompt_frames", job_id)
    os.makedirs(frames_dir, exist_ok=True)

    for index, frame_url in enumerate(frame_urls):
        response = httpx.get(frame_url, timeout=60.0)
        response.raise_for_status()
        match = re.search(r"frame[_-]?(\d+)", frame_url, re.IGNORECASE)
        frame_number = match.group(1) if match else str(index)
        frame_key = f"frame_{frame_number}"
        filename = f"{frame_key}.jpg"
        frame_path = os.path.abspath(os.path.join(frames_dir, filename))
        with open(frame_path, "wb") as f:
            f.write(response.content)
        frame_images[frame_key] = frame_path
    return frame_images


def _download_frame_images_from_job_zip(job_id: str, backend_base_url: str) -> dict[str, str]:
    if not job_id or not backend_base_url:
        return {}

    frames_dir = os.path.join("cat", "static", "video_prompt_frames", job_id)
    os.makedirs(frames_dir, exist_ok=True)

    download_url = f"{backend_base_url}/api/jobs/{job_id}/download"
    response = httpx.get(download_url, timeout=120.0)
    response.raise_for_status()

    frame_images: dict[str, str] = {}
    with zipfile.ZipFile(BytesIO(response.content)) as job_zip:
        for entry in job_zip.infolist():
            if entry.is_dir():
                continue
            if not entry.filename.lower().startswith("frames/"):
                continue

            basename = os.path.basename(entry.filename)
            match = re.search(r"frame[_-]?(\d+)", basename, re.IGNORECASE)
            if not match:
                continue

            frame_number = match.group(1)
            frame_key = f"frame_{frame_number}"
            target_path = os.path.abspath(os.path.join(frames_dir, basename))
            with job_zip.open(entry) as source, open(target_path, "wb") as dest:
                dest.write(source.read())
            frame_images[frame_key] = target_path

    return frame_images


def _insert_frame_image_before_paragraph(paragraph, frame_path: str):
    image_paragraph = paragraph.insert_paragraph_before()
    image_run = image_paragraph.add_run()
    image_run.add_picture(frame_path, width=Inches(5.8))


def _replace_frame_placeholders_with_images(doc, frame_images: dict[str, str]) -> tuple[int, int]:
    replacements = 0
    placeholders_found = 0
    for paragraph in list(doc.paragraphs):
        original_text = paragraph.text or ""
        paragraph_text = original_text.strip()
        matches = re.findall(r"\[frame_\d+\]", paragraph_text.lower())
        if not matches:
            continue

        placeholders_found += len(matches)
        handled_any = False
        for match in matches:
            frame_key = match.strip("[]")
            frame_path = frame_images.get(frame_key)
            if not frame_path or not os.path.exists(frame_path):
                continue

            _insert_frame_image_before_paragraph(paragraph, frame_path)
            handled_any = True
            replacements += 1

        if handled_any:
            cleaned_text = re.sub(r"\s*\[frame_\d+\]\s*", " ", original_text, flags=re.IGNORECASE)
            paragraph.text = cleaned_text.strip()
    return placeholders_found, replacements


def _start_video_upload_session(cat) -> dict:
    session = create_session(str(getattr(cat, "user_id", "") or ""))

    upload_url = f"{UPLOAD_PAGE_URL}?session_id={session['session_id']}"
    upload_link = (
        f'<a href="{upload_url}" target="_blank" rel="noopener noreferrer">'
        "Apri Video Analyzer Upload"
        "</a>"
    )

    cat.send_ws_message(
        "Se vuoi inviare un video, caricalo qui: "
        f"{upload_link}",
        "chat",
    )

    deadline = time.time() + WAIT_TIMEOUT_SECONDS
    while time.time() < deadline:
        current = get_session(session["session_id"]) or {}
        status = current.get("status")

        result = current.get("result")

        if status == "completed" and result not in (None, ""):
            delete_session(session["session_id"])
            return {
                "status": "completed",
                "result": clean_result_string(result),
                "upload_url": upload_url,
            }

        if status == "failed":
            delete_session(session["session_id"])
            return {
                "status": "failed",
                "error": current.get("error") or "Errore durante l'analisi del video.",
                "upload_url": upload_url,
            }

        if status == "cancelled":
            delete_session(session["session_id"])
            return {
                "status": "cancelled",
                "error": "Caricamento video analyzer interrotto.",
                "upload_url": upload_url,
            }

        time.sleep(POLL_INTERVAL_SECONDS)

    return {
        "status": "timeout",
        "error": (
            "Timeout in attesa del caricamento video. "
            f"Puoi completare l'upload qui: {upload_url}"
        ),
        "upload_url": upload_url,
    }


def _chunk_video_memory_text(cat, result_text: str, source: str, metadata: dict):
    """Divide il testo della memoria video usando il chunker standard del Rabbit Hole."""
    docs = cat.rabbit_hole.string_to_docs(
        cat=cat,
        file_bytes=result_text.encode("utf-8"),
        source=source,
        content_type="text/plain",
    )

    for doc in docs:
        current_metadata = dict(getattr(doc, "metadata", {}) or {})
        current_metadata.update(metadata)
        current_metadata["source"] = source
        doc.metadata = current_metadata

    return docs


def _remember_video_result(cat, result_text: str, original_command: str) -> str:
    title_prompt = f"""Genera un titolo breve e descrittivo (massimo 5 parole) per questa analisi video:

\"{result_text}\"

Rispondi SOLO con il titolo, senza punteggiatura finale."""

    generated_title = cat.llm(title_prompt).strip().strip("\"'.,;:!?")
    if not generated_title:
        generated_title = "Analisi video"

    safe_filename = re.sub(r"[^\w\s-]", "", generated_title)
    safe_filename = re.sub(r"[-\s]+", "_", safe_filename).strip("_") or "analisi_video"
    timestamp = int(time.time())
    filename = f"{safe_filename}_{timestamp}.txt"

    static_dir = "cat/static"
    os.makedirs(static_dir, exist_ok=True)
    filepath = os.path.join(static_dir, filename)

    file_content = f"""Titolo: {generated_title}
Data creazione: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
Tipo: Memoria Analisi Video
Utente: {getattr(cat, "user_id", "")}

---

{result_text}
"""

    with open(filepath, "w", encoding="utf-8") as f:
        f.write(file_content)

    created_at = time.time()
    video_memory_metadata = {
        "source": generated_title,
        "type": "video_analysis_memory",
        "user_input": True,
        "original_command": original_command,
        "created_at": created_at,
        "file_path": filepath,
    }

    docs = _chunk_video_memory_text(
        cat=cat,
        result_text=result_text,
        source=generated_title,
        metadata=video_memory_metadata,
    )

    log.info(f"[Video Analyzer Remember Hook] Memoria video suddivisa in {len(docs)} chunk")

    cat.rabbit_hole.store_documents(
        cat=cat,
        docs=docs,
        source=generated_title,
        metadata=video_memory_metadata,
    )

    cat.send_ws_message(
        f"💾 {generated_title}: {result_text[:50]}...",
        msg_type="notification",
    )
    cat.send_ws_message(
        f'📄 File creato: <a href="/static/{filename}?v={timestamp}" download="{filename}">Scarica {generated_title}.txt</a>',
        "chat",
    )

    return generated_title



# - usa questo schema esatto:
#   [frame_2]
#   16. Identificare ... .
# - non usare mai formati come `16. [frame_2] Identificare...` oppure `• [frame_2] ...`;
# - usa esclusivamente tag realmente presenti nel contenuto fornito, ad esempio `[frame_2]`, `[frame_6]`, `[frame_8]`;
# - se il contenuto supporta un riferimento visivo, non omettere il tag frame;

def _build_video_procedure_prompt(result_text: str, tool_input: str) -> str:
    user_requirement = (tool_input or "").strip() or (
        "Trasforma l'analisi video in una procedura operativa standard, chiara e riutilizzabile."
    )
    return f"""Sei un redattore tecnico esperto nella stesura di procedure operative.

Trasforma il seguente risultato di analisi video in una procedura reale, pulita e direttamente eseguibile da un operatore esterno.

Regole obbligatorie:
- descrivi dettagliatamente tutte le informazioni utili all'operatore;
- usa solo informazioni supportate dal contenuto fornito;
- non inventare azioni, click, schermate o valori non presenti;
- non citare video, timestamp, JSON, trascrizione, osservazioni o evidenze;
- non scrivere frasi come "nel video", "si osserva", "dalla trascrizione";
- converti il contenuto in istruzioni operative dirette;
- scrivi in italiano tecnico, chiaro, sintetico e professionale.

Formato obbligatorio:
1. Titolo
2. Scopo
3. Procedura operativa

Regole per la sezione "Procedura operativa":
- usa passi numerati;
- ogni passo deve descrivere un'azione concreta;
- se utile, organizza i passi in fasi;
- l'output deve sembrare una SOP, non un'analisi.
- quando un passaggio richiede un riferimento visivo utile, inserisci il tag frame su una riga tutta sua, separata dal testo;
- il tag frame non deve mai stare nella stessa riga del numero passo, del bullet o della frase operativa;
- il tag frame deve comparire esattamente da solo nel formato `[frame_x]`;
- subito sotto il tag frame deve esserci l'istruzione operativa a cui si riferisce;
- non è necessario usare tutti i frame;
- non aggiungere commenti descrittivi sul frame: usa solo il tag essenziale `[frame_x]`.

Personalizzazione richiesta dall'utente:
{user_requirement}

Contenuto di partenza:
{result_text}
"""


# @tool(
#     return_direct=False,
#     examples=[
#         "analizza il video",
#         "voglio analizzare un video",
#         "apri l'upload per analizzare un video",
#         "carico un filmato da analizzare",
#         "fammi la trascrizione di un video",
#         "analizza questo video e dammi il json strutturato",
#         "usa il video analyzer come passo preliminare",
#     ],
# )
# def video_analyzer(tool_input: str, cat):
#     """
#     Use this tool whenever the agent needs to analyze a video.
#     Use it even if the user has not uploaded the file yet: this tool sends the upload link in chat and waits for the result.
#     The input is the desired analysis task. 
#     The output is a structured JSON with the cleaned video analysis result.
#     This tool is typically a preliminary step used to prepare input for a following tool.
#     """

@tool(
    return_direct=True,
    examples=[
        "analizza il video",
        "voglio analizzare un video",
        "apri l'upload per analizzare un video",
        "carico un filmato da analizzare",
        "fammi la trascrizione di un video",
        "analizza questo video e trasformalo in una procedura",
        "usa il video analyzer per ricavare una procedura operativa",
    ],
)
def video_analyzer(tool_input: str, cat):
    """
    Use this tool whenever the agent needs to analyze a video.
    Use it even if the user has not uploaded the file yet: this tool sends the upload link in chat and waits for the result.
    The input is the desired transformation to apply to the analyzed video.
    The tool first obtains the cleaned structured video result, then converts it into a personalized procedure influenced by the user request.
    The final output is the confirmation that the procedure has been created.
    """
    session_result = _start_video_upload_session(cat)
    if session_result["status"] != "completed":
        return session_result["error"]

    result_text = session_result["result"]
    prompt_personalizzato = _build_video_procedure_prompt(result_text, tool_input)
    procedura_finale = cat.llm(prompt_personalizzato)
    doc = load_template_document(cat)
    parse_markdown_to_word(doc, procedura_finale)

    job_id = _extract_test_job_id(result_text)
    backend_base_url = _get_backend_base_url(cat)
    if job_id and backend_base_url:
        try:
            frame_urls = _fetch_backend_frame_urls(job_id, backend_base_url)
            if frame_urls:
                try:
                    frame_images = _download_frame_images_from_job_zip(job_id, backend_base_url)
                except Exception:
                    frame_images = _download_frame_images_to_static(frame_urls, job_id)
                _replace_frame_placeholders_with_images(doc, frame_images)
        except Exception as frame_error:
            log.warning(f"[Video Analyzer Tool] Recupero frame fallito: {frame_error}")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"video_procedure_{timestamp}.docx"
    static_dir = "cat/static"
    os.makedirs(static_dir, exist_ok=True)
    file_path = os.path.join(static_dir, filename)
    doc.save(file_path)
    download_url = f"/static/{filename}?v={timestamp}"

    #TODO: CANCELLARE JOBID DATA DAL BACKEND

    # cat.send_ws_message(
    #     f"Tool_input: {tool_input}",
    #     "chat",
    # )

    # cat.send_ws_message(
    #     f"Prompt personalizzato: {prompt_personalizzato}",
    #     "chat",
    # )

    # cat.send_ws_message(
    #     f"procedura finale:\n{procedura_finale}",
    #     "chat",
    # )

    # cat.send_ws_message(
    #     f"job_id: {job_id}",
    #     "chat",
    # )

    # cat.send_ws_message(
    #     f'📄 File Word creato: <a href="{download_url}" download="{filename}">Scarica {filename}</a>',
    #     "chat",
    # )

    # cat.send_ws_message(
    #             f'📄 File Word creato: <a href="{download_url}" download="{filename}">Scarica {filename}</a>',
    #             "chat",
    #         )

    _delete_backend_job_if_possible(cat, job_id)
    return f'📄 File Word creato: <a href="{download_url}" download="{filename}">Scarica {filename}</a>'


@hook
def fast_reply(fast_reply, cat):
    settings = cat.mad_hatter.get_plugin().load_settings()
    tool_key = settings.get("remember_tool_name", "ricorda il video: ...il video...")

    if not _is_tool_enabled_for_user(cat, tool_key):
        return fast_reply

    user_message = cat.working_memory.get("user_message_json", {}).get("text", "")
    if not re.match(r"^ricorda\s+il\s+video\s*$", user_message.strip(), re.IGNORECASE):
        return fast_reply

    try:
        session_result = _start_video_upload_session(cat)
        if session_result["status"] != "completed":
            return {"output": session_result["error"]}

        result_text = str(session_result["result"])
        title = _remember_video_result(
            cat,
            result_text=result_text,
            original_command="ricorda il video",
        )
        job_id = _extract_test_job_id(result_text)
        _delete_backend_job_if_possible(cat, job_id)
        return {
            "output": f"""✅ **Ho memorizzato il video:**\n\n*{title}"""
        }
    except Exception as e:
        log.error(f"[Video Analyzer Remember Hook] Errore durante la memorizzazione: {str(e)}")
        return {
            "output": f"❌ Si è verificato un errore durante la memorizzazione del video:\n\n`{str(e)}`"
        }


## PROMPT DEBUG
# str_json_text="""
# {
#   "job_id": "d500613d-bba3-4725-921c-7cdc715298b9",
#   "analysis": {
#     "frame_analyses": [
#       {
#         "response": "```\nFrame 0 (116.16s)\n\nSetting/Scene:\n- Screen capture of an industrial automation/PLC programming IDE (CODESYS-like) on Windows.\n- Left panel shows a “Targets” tree with “PC-Simulator” expanded and multiple IEC tasks listed (e.g., IEC_Display_Task, IEC_Exception_Task, IEC_FS_TASK, IEC_MainTask, IEC_MBM_LineTask, IEC_Supervisor_Task variants, IEC_UITask). A “PLC” entry appears below.\n\nAction/Movement:\n- No visible cursor movement or UI interaction in this frame; the interface is static.\n\nNew Information:\n- Center “Output” pane displays build messages/warnings about library versions (multiple lines indicating referenced libraries have lower versions than latest used in the project).\n- A line indicates the build completed successfully (e.g., “Build completed without errors” and elapsed time shown).\n- Bottom pane shows a “Global variables” table for a project (name visible as something like “SWHQOORTB00MM00”), with variable names and types (e.g., REAL/BOOL/WORD/UDINT/ARRAY), including entries such as “AbsHumSetOffset” (REAL) and several “Access_*” BOOL variables.\n- Status bar at the bottom indicates connection/synchronization to “PC-Simulator” (green bar).\n\nContinuity Points:\n- This is the starting frame; establishes that the video is focused on PLC/automation project build output and"
#       },
#       {
#         "response": "```\nFrame 1 (146.39s)\n\nSetting/Scene (if changed from previous):\n- A web browser window is now in the foreground over the PLC IDE, showing “STone - Web simulator” (URL visible as 127.0.0.1:43051/...).\n- The PLC IDE remains visible in the background (Targets tree on the left; Output and Global variables panes partially visible). Bottom status bar still shows a green “Connected to PC-Simulator, Synchronized” strip.\n\nAction/Movement:\n- No cursor movement is visible; the scene appears static with the simulator open and displaying a device UI.\n\nNew Information:\n- The STone simulator shows a device screen with a “Language” menu: “Language: EN” and instructions “UP/DOWN to change” and “ENTER to confirm.”\n- Left side of the simulator has a menu list: “Toggle compact layout,” “Toggle I/O visibility,” “Alarms,” “Screenshots,” and “Credits,” plus a “Collapse menu” option at the bottom.\n- Additional simulator sections are visible below: “MULTI KEYS SEND” (radio options “Key codes” and “Custom key code,” with fields “Key code-1” and “Key code-2,” and a disabled-looking “SEND” button), “SEVEN SEGMENTS,” and “LEDS (24)” with several green indicator dots and a camera icon button.\n\nContinuity Points:\n- Continues from the"
#       },
#       {
#         "response": "```\nFrame 2 (146.87s)\n\nSetting/Scene (if changed from previous):\n- The web browser with “STone - Web simulator” remains in the foreground. The simulator layout is more fully visible, showing additional I/O panels below the device screen.\n\nAction/Movement:\n- No visible cursor movement or interaction; the screen appears static.\n\nNew Information:\n- The simulated device display still shows the “Language” screen (“Language: EN” with “UP/DOWN to change” and “ENTER to confirm”), with a small time indicator at the top right of the device display (appears as “4:9”/similar).\n- Left sidebar menu remains: “Toggle compact layout,” “Toggle I/O visibility,” “Alarms,” “Screenshots,” “Credits,” and “Collapse menu.”\n- Additional simulator panels are clearly visible:\n  - “DIGITAL INPUT (24/24)” with multiple toggle switches labeled ID1–ID24 (blue toggles shown).\n  - “SEVEN SEGMENTS” section with “UNIVERSAL (10)” showing circular dial-like controls labeled U1–U10 with small numeric values near each.\n  - “LEDS (24)” grid labeled 0–23 with green indicator dots.\n  - “DIGITAL OUTPUT (29/29)” panel visible on the right with outputs labeled DO1, DO2, etc., showing green indicator dots.\n- A small floating camera icon/button appears near the lower-right"
#       },
#       {
#         "response": "```\nFrame 3 (148.79s)\n\nSetting/Scene (if changed from previous):\n- The STone - Web simulator remains in the foreground in a browser (127.0.0.1:43051/...).\n- View is slightly scrolled/positioned so the device screen is only partially visible at the top; the I/O panels dominate the view.\n\nAction/Movement:\n- No visible interaction; all controls and indicators appear static.\n\nNew Information:\n- Left sidebar shows: “Toggle compact layout,” “Toggle I/O visibility,” “Alarms,” “Screenshots,” “Credits,” and “Collapse menu” at the bottom.\n- “MULTI KEYS SEND” panel is visible with “Key codes” selected; fields labeled “Key code-1” and “Key code-2” are present; “SEND” button appears disabled/greyed out.\n- “DIGITAL INPUT (24/24)” panel shows ID1–ID24 with blue toggle switches (all appear in the same on-position as prior frames).\n- “SEVEN SEGMENTS” → “UNIVERSAL (10)” shows U1–U10 as blue circular dials with small numeric readouts near each (values around 0.01, -0.03, -0.02, etc.).\n- “LEDS (24)” grid labeled 0–23 shows green indicator dots lit.\n- “DIGITAL OUTPUT (29/29)” panel shows DO1"
#       },
#       {
#         "response": "```\nFrame 4 (149.27s)\n\nSetting/Scene (if changed from previous):\n- STone - Web simulator remains in the foreground in a browser (127.0.0.1:43051/...).\n- View is positioned so only the lower portion of the simulated device screen is visible at the top (showing the instruction text area), while the I/O panels fill most of the page.\n\nAction/Movement:\n- No visible cursor movement or interaction; all toggles/dials/LED indicators appear unchanged and static.\n\nNew Information:\n- The top of the device display shows partial text: “UP/DOWN to change” and “ENTER to confirm” (the “Language” header itself is mostly off-screen).\n- “MULTI KEYS SEND” panel shows “Key codes” selected, with two input fields labeled “Key code-1” and “Key code-2”; “SEND” button remains greyed out/disabled.\n- “DIGITAL INPUT (24/24)” panel shows ID1–ID24 with blue toggle switches (all appear in the same on-position as earlier frames).\n- “SEVEN SEGMENTS” → “UNIVERSAL (10)” shows U1–U10 blue circular dials with small numeric values (e.g., around 0.01, -0.01, -0.02, -0.03).\n- “LEDS (24)” grid labeled 0–23 shows green indicator dots"
#       },
#       {
#         "response": "```\nFrame 5 (149.75s)\n\nSetting/Scene (if changed from previous):\n- The STone - Web simulator remains in the foreground in a browser (127.0.0.1:43051/...).\n- The page is positioned slightly higher than the prior frame, so more of the simulated device screen is visible at the top.\n\nAction/Movement:\n- No visible cursor movement or interaction; all UI elements appear static.\n\nNew Information:\n- The simulated device display now shows more of the header area: “Language:” is visible with “EN” on the right; the top-right time indicator reads approximately “4:19.”\n- The instruction text remains visible: “UP/DOWN to change” and “ENTER to confirm.”\n- Left sidebar menu remains visible: “Toggle compact layout,” “Toggle I/O visibility,” “Alarms,” “Screenshots,” “Credits,” with “Collapse menu” at the bottom.\n- “MULTI KEYS SEND” panel remains with “Key codes” selected; fields “Key code-1” and “Key code-2” are present; “SEND” button remains disabled/greyed out.\n- “SEVEN SEGMENTS” → “UNIVERSAL (10)” dials U1–U10 remain visible with small numeric values (e.g., U1 ~0.01, U2 ~-0.01, U3 ~-0.02, U4 ~-0.00,"
#       },
#       {
#         "response": "```\nFrame 6 (160.31s)\n\nSetting/Scene (if changed from previous):\n- The STone - Web simulator remains in the foreground in a browser (127.0.0.1:43051/...).\n- The simulated device screen content has changed from the prior “Language” menu to a status/measurement screen labeled “Info - Circuit 1.”\n\nAction/Movement:\n- No visible cursor movement or interaction; panels and indicators appear static.\n\nNew Information:\n- The device display now shows “Info - Circuit 1” at the top, with multiple readouts and labels including “Req:” and “Run:” (both appear to show 0%).\n- Several numeric values are visible on the device screen, including “0.0bar” and temperatures around “-51.x°C” (e.g., -51.5°C / -51.4°C), plus “0.0°C” entries.\n- A “Status:” line is visible (text appears as “VLV NO GO” or similar), and a small label near the lower-left of the device screen reads “LoSCP” (or similar).\n- The simulator’s left sidebar remains: “Toggle compact layout,” “Toggle I/O visibility,” “Alarms,” “Screenshots,” “Credits,” and “Collapse menu.”\n- The “MULTI KEYS SEND” panel still shows “Key codes” selected with “Key code-1” and “Key code-2"
#       },
#       {
#         "response": "```\nFrame 7 (160.79s)\n\nSetting/Scene (if changed from previous):\n- STone - Web simulator remains open in a browser (127.0.0.1:43051/...); view is scrolled so only the lower portion of the simulated device screen is visible at the top, with I/O panels filling most of the page.\n\nAction/Movement:\n- No visible cursor movement or interaction; controls and indicators appear unchanged.\n\nNew Information:\n- The device screen portion visible still corresponds to the “Info - Circuit 1” status/measurement screen; readable lines include “Status: VLV NO GO” and values such as “0.0bar” and temperatures around “-51.4°C / -51.5°C,” plus “0.0°C.”\n- “MULTI KEYS SEND” panel shows “Key codes” selected; fields “Key code-1” and “Key code-2” are present; “SEND” button remains disabled/greyed out.\n- “SEVEN SEGMENTS” → “UNIVERSAL (10)” shows dials U1–U10 with small numeric values (e.g., U1 ~0.01, U2 ~-0.02, U3 ~-0.02, U4 ~-0.02, U5 ~0.03, U6 ~0.02, U7 ~-0.00, U8 ~-0.01,"
#       },
#       {
#         "response": "```\nFrame 8 (161.27s)\n\nSetting/Scene (if changed from previous):\n- STone - Web simulator remains in the foreground in a browser (127.0.0.1:43051/...).\n- View is positioned so the simulated device screen (“Info - Circuit 1”) is visible at the top center, with I/O panels below (Multi Keys Send, Seven Segments/Universal dials, LEDs, Digital Input, Digital Output).\n\nAction/Movement:\n- A mouse cursor is visible over the “SEVEN SEGMENTS” → “UNIVERSAL (10)” area, hovering on/near one of the blue circular dials (around U7/U8 region).\n- A tooltip appears over the dial, indicating hover interaction; no dial rotation or value change is visible.\n- No other UI elements change; toggles and LEDs remain steady.\n\nNew Information:\n- Tooltip text appears over the Universal dial: \n  - “Rotate the knob to change the value.”\n  - “Use SHIFT to increase/decrease the value with more precision.”\n  - “Use CTRL to increase/decrease the value faster.”\n- The device screen continues to show “Info - Circuit 1” with “Req: 0%” and “Run: 0%,” “0.0bar,” temperatures around “-51.5°C / -51.4°C,” and “Status: VLV NO GO,” plus the “LoSCP” label"
#       },
#       {
#         "response": "```\nFrame 9 (177.11s)\n\nSetting/Scene (if changed from previous):\n- The view returns to the PLC programming IDE (CODESYS-like) as the main background window on Windows.\n- A smaller web browser window/tab group is in the foreground, showing “STone - Web simulator” (127.0.0.1:43051/…); it is not full-screen and sits over the IDE.\n\nAction/Movement:\n- No visible cursor movement or interaction in this frame; all windows appear static.\n- The simulator remains open but appears repositioned/resized compared to the prior full-page simulator view.\n\nNew Information:\n- The IDE’s left “Targets” tree is visible again with “PC-Simulator” expanded and multiple IEC tasks listed.\n- The IDE “Output” pane shows the same style of build/library version warnings and “Post-build completed”/build completion lines.\n- The IDE bottom status bar shows a green strip indicating “Connected to PC-Simulator, Synchronized.”\n- In the foreground STone simulator device screen, the display remains on “Info - Circuit 1,” showing “Req: 0%” and “Run: 0%,” “0.0bar,” temperatures around “-51.x°C,” and “Status: VLV NO GO,” with the “LoSCP” label visible.\n- The simulator’s left sidebar menu is visible (e.g., “Toggle compact layout,” “Toggle I/O"
#       },
#       {
#         "response": "```\nFrame 10 (177.59s)\n\nSetting/Scene (if changed from previous):\n- The PLC programming IDE (CODESYS-like) remains the main background window with the “Targets” tree on the left and “Output”/“Global variables” panes visible.\n- A smaller browser window stays in the foreground showing “STone - Web simulator” (127.0.0.1:43051/...), centered over the IDE.\n\nAction/Movement:\n- No visible cursor movement or interaction; both the IDE and simulator appear static.\n\nNew Information:\n- In the STone simulator device screen (“Info - Circuit 1”), the readouts now include additional non-zero values compared to earlier frames:\n  - “Req: 0%” and “Run: 0%” remain visible.\n  - A pressure value “1.6bar” appears on a lower line (previously the screen showed “0.0bar”).\n  - A temperature value around “-31.0°C” appears on the lower portion of the display (previously temperatures were around “-51.x°C”).\n  - “Status: VLV NO GO” remains visible.\n- The simulator left sidebar still lists: “Toggle compact layout,” “Toggle I/O visibility,” “Alarms,” “Screenshots,” and “Credits,” with “Collapse menu” at the bottom.\n- “MULTI KEYS SEND” panel remains with “Key codes” selected and"
#       }
#     ],
#     "transcript": " Okay, so here we are on Stone and now let's describe the procedure to launch the program simulator.  I close the windows here just for simplicity.  Okay, here we have in Solution Explorer, I go to SWH here on the program,  ah no, here, full build, right-right, full build, on the base program, SWH, Q00, etc.  Then in Output, there is the log of the build, build completed without errors.  Okay, now we have to launch the simulator and connect to the target.  So if you have the target down here, you press here and go here.  If you don't have it there, you have to go to View, Target, which always leads here.  Now here we have PC Simulator, this allows you to simulate the PLC on this computer, on this PC.  If I press the right button here and do Connect to Target, the red X comes out,  the red X comes out because the file is not compiled yet.  Let's do full download, so right-click on PC Simulator, full download,  and we have seen that it has become green.  This means that the project compiled on PC Simulator is equal to your source.  Now we go here on the right and we press the button here with a screen,  which then opens a web page where there are all the web simulators,  where there are, for example, all the values, here you can simulate the PLC variables.",
#     "video_description": {
#       "response": "```\nVIDEO SUMMARY\nDuration: ~177s (based on visible timestamps in frame notes)\n\nThe video opens on a Windows desktop screen capture of an industrial automation/PLC programming IDE (CODESYS-like). The “Targets” tree on the left shows a “PC-Simulator” target with multiple IEC tasks listed. In the center “Output” pane, build messages include library version warnings and a clear “build completed without errors” result. A “Global variables” table is visible at the bottom, and the status bar indicates the system is connected/synchronized to the PC-Simulator.\n\nThe narrator explains the procedure to launch the program simulator. They describe going to the project in the Solution Explorer and running a full build, then checking the Output log to confirm the build completes successfully. Next, they explain how to open the Targets view (via the target panel or through View → Target) and select “PC Simulator” to simulate the PLC on the local computer. They mention that attempting “Connect to Target” can show a red X if the program is not compiled/downloaded yet, and then they perform (or describe performing) a “Full download” on the PC-Simulator target so it turns green, indicating the compiled project on the simulator matches the source.\n\nAfter this, a web browser window comes to the foreground showing “STone - Web simulator” hosted locally (127.0.0.1). Initially, the simulated device screen displays a “Language” menu (“Language: EN” with instructions to use UP/DOWN and ENTER). The simulator interface includes a left sidebar (compact layout toggle, I/O visibility, alarms, screenshots, credits) and multiple I/O panels below the device screen: a “MULTI KEYS SEND” section (with key code fields and a disabled-looking SEND button), “DIGITAL INPUT (24/24)” toggles, “SEVEN SEGMENTS” universal dial controls (U1–U10), a grid of “LEDS (24)” lit green, and “DIGITAL OUTPUT (29/29)” indicators. The page view shifts slightly (as if scrolled) to emphasize the I/O panels, and at one point the mouse hovers over a universal dial, triggering a tooltip explaining how to rotate the knob and use SHIFT/CTRL for precision/speed.\n\nThe simulated device screen then changes from the language menu to an operational status page titled “Info - Circuit 1,” showing readouts such as “Req: 0%,” “Run: 0%,” pressure and temperature values (including very low temperatures around -51°C), and a status line reading “VLV NO GO,” with a small label near the bottom (e.g., “LoSCP”). Near the end, the view returns to the PLC IDE with the STone simulator now in a smaller foreground window. The “Info - Circuit 1” screen remains visible, and some values update compared to earlier (e.g., pressure showing around 1.6 bar and a temperature around -31°C), while the IDE behind still shows the targets tree, build output, and a green connected/synchronized status.\n\nClosing observations: The clip ends with the PLC project successfully built and downloaded to the PC-Simulator (green/connected state), and the STone web simulator open to visualize and simulate PLC-related values and I/O through the “Info - Circuit 1” screen and the interactive input/output panels.\n\nNote: This summary is based on direct observation of the first frame combined with detailed notes from subsequent frames.\n```"
#     }
#   }
# }
# """

# tool_input_test="fai una procedura passo passo"


# @hook
# def fast_reply(fast_reply, cat):
#     user_message = cat.working_memory.get("user_message_json", {}).get("text", "")
#     if not re.match(r"^test\s+prompt\s+video\s*$", user_message.strip(), re.IGNORECASE):
#         return fast_reply
#     ...
#     try:
#         result_text = (
#             str_json_text.replace("“", '"')
#             .replace("”", '"')
#             .replace("‘", "'")
#             .replace("’", "'")
#         )
#         prompt_personalizzato = _build_video_procedure_prompt(result_text, tool_input_test)

#         cat.send_ws_message(f"Tool_input: {tool_input_test}", "chat")
#         cat.send_ws_message(f"Prompt personalizzato: {prompt_personalizzato}", "chat")

#         procedure_result = cat.llm(prompt_personalizzato)
#         cat.send_ws_message(f"procedura finale:\n{procedure_result}", "chat")

#         doc = load_template_document(cat)
#         parse_markdown_to_word(doc, procedure_result)
#         job_id = _extract_test_job_id(result_text)
#         cat.send_ws_message(f"job_id: {job_id}", "chat")
#         backend_base_url = _get_backend_base_url(cat)
#         if job_id and backend_base_url:
#             try:
#                 cat.send_ws_message(f"Recupero frame backend per job_id: {job_id}", "chat")
#                 frame_urls = _fetch_backend_frame_urls(job_id, backend_base_url)
#                 if frame_urls:
#                     cat.send_ws_message(
#                         f"Trovati {len(frame_urls)} frame dal backend. Li inserisco nel Word dove trova [frame_x].",
#                         "chat",
#                     )
#                     try:
#                         frame_images = _download_frame_images_from_job_zip(job_id, backend_base_url)
#                         cat.send_ws_message(
#                             "Frame recuperati dallo ZIP del job backend.",
#                             "chat",
#                         )
#                     except Exception as zip_error:
#                         log.warning(f"[Video Analyzer Prompt Test] Recupero frame da ZIP fallito: {zip_error}")
#                         cat.send_ws_message(
#                             f"⚠️ Recupero frame da ZIP fallito, provo il download diretto dei file: {zip_error}",
#                             "chat",
#                         )
#                         frame_images = _download_frame_images_to_static(frame_urls, job_id)
#                     cat.send_ws_message(
#                         f"Frame salvati in appoggio locale: cat/static/video_prompt_frames/{job_id}",
#                         "chat",
#                     )
#                     placeholders_found, inserted_frames = _replace_frame_placeholders_with_images(doc, frame_images)
#                     cat.send_ws_message(
#                         f"Placeholder frame trovati nel testo: {placeholders_found}",
#                         "chat",
#                     )
#                     cat.send_ws_message(
#                         f"Frame inseriti nel Word: {inserted_frames}",
#                         "chat",
#                     )
#                 else:
#                     cat.send_ws_message("Nessun frame disponibile dal backend per questo job.", "chat")
#             except Exception as frame_error:
#                 log.warning(f"[Video Analyzer Prompt Test] Recupero frame fallito: {frame_error}")
#                 cat.send_ws_message(
#                     f"⚠️ Recupero frame non riuscito, creo comunque il Word senza immagini: {frame_error}",
#                     "chat",
#                 )

#         timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
#         filename = f"video_prompt_test_{timestamp}.docx"
#         static_dir = "cat/static"
#         os.makedirs(static_dir, exist_ok=True)
#         file_path = os.path.join(static_dir, filename)
#         doc.save(file_path)

#         download_url = f"/static/{filename}?v={timestamp}"
#         cat.send_ws_message(
#             f'📄 File Word creato: <a href="{download_url}" download="{filename}">Scarica {filename}</a>',
#             "chat",
#         )
#         return {
#             "output": f'Procedura Word pronta: <a href="{download_url}" download="{filename}">Scarica il file</a>'
#         }
#     except Exception as e:
#         log.error(f"[Video Analyzer Prompt Test] Errore durante il test prompt: {str(e)}")
#         return {"output": f"❌ Errore durante il test prompt:\n\n`{str(e)}`"}


# Output_ottenuto = """
#     procedura finale:

#     Titolo
#     Utilizzo del simulatore STone (Web simulator) con target “PC-Simulator” nel progetto PLC

#     Scopo
#     Eseguire build e download del progetto sul target “PC-Simulator”, avviare il simulatore STone (interfaccia web locale) e utilizzare le funzioni principali di visualizzazione/simulazione I/O e schermate dispositivo.

#     Prerequisiti

#     PC Windows con ambiente di sviluppo PLC (IDE con pannelli “Targets”, “Output”, “Global variables”).
#     Progetto PLC disponibile (es. progetto con nome visibile “SWHQOORTB00MM00”).
#     Target “PC-Simulator” disponibile nella vista “Targets”.
#     Browser con accesso al simulatore locale “STone - Web simulator” su indirizzo 127.0.0.1 (porta indicata nell’URL).
#     Permessi per eseguire build e download sul target.
#     Procedura operativa
#     Fase A — Build del progetto

#     Nel progetto PLC, eseguire Full build del programma di base (es. ramo/progetto “SWH …”).
#     Aprire il pannello Output e verificare la presenza del messaggio Build completed without errors.
#     Fase B — Apertura vista Target e selezione “PC-Simulator”

#     Se la sezione target non è già visibile, aprire il menu View e selezionare Target.
#     Nella vista Targets, individuare ed espandere PC-Simulator.
#     Fase C — Connessione e download sul target

#     Su PC-Simulator, aprire il menu contestuale (tasto destro).
#     (Opzionale) Se si seleziona Connect to Target e compare una X rossa, procedere con il download completo come indicato nei passi successivi.
#     Su PC-Simulator, eseguire Full download dal menu contestuale.
#     Verificare che PC-Simulator risulti in stato verde, a conferma che il progetto compilato sul simulatore è allineato ai sorgenti.
#     Verificare nella barra di stato dell’IDE lo stato Connected to PC-Simulator, Synchronized.
#     Fase D — Avvio del simulatore STone (Web simulator)
#     [Frame 1]

#     Nell’IDE, premere il pulsante con icona schermo per aprire la pagina web dei simulatori STone.
#     Nel browser, verificare l’apertura di STone - Web simulator su indirizzo locale 127.0.0.1.
#     Fase E — Utilizzo delle funzioni principali del simulatore
#     [Frame 2]

#     Nella barra laterale del simulatore, utilizzare le voci disponibili in base alla necessità:
#     Toggle compact layout
#     Toggle I/O visibility
#     Alarms
#     Screenshots
#     Credits
#     Collapse menu
#     Nella sezione DIGITAL INPUT (24/24), utilizzare i toggle ID1–ID24 per simulare gli ingressi digitali.
#     Nella sezione DIGITAL OUTPUT (29/29), consultare gli indicatori DO1… per monitorare lo stato delle uscite digitali.
#     Nella sezione LEDS (24), consultare la griglia LED (0–23) per il monitoraggio visivo degli stati.
#     Nella sezione SEVEN SEGMENTS → UNIVERSAL (10), utilizzare le manopole U1–U10 per variare i valori associati.
#     [Frame 8]

#     Per modificare un valore con le manopole UNIVERSAL (10):
#     Ruotare la manopola per cambiare il valore.
#     Usare SHIFT per aumentare/diminuire con maggiore precisione.
#     Usare CTRL per aumentare/diminuire più velocemente.
#     Nella sezione MULTI KEYS SEND, mantenere selezionata l’opzione Key codes e utilizzare i campi Key code-1 e Key code-2 quando necessario (eseguire l’invio solo se il pulsante SEND risulta attivo).
#     Fase F — Monitoraggio schermate dispositivo
#     [Frame 6]

#     Sul display del dispositivo simulato, monitorare le schermate operative disponibili, ad esempio Info - Circuit 1, verificando i principali indicatori (es. Req, Run, pressione in bar, temperature in °C, e riga Status).
#     Fase G — Arresto/chiusura

#     Per terminare la sessione di simulazione, chiudere la finestra/tab del browser STone - Web simulator oppure ridurla se si necessita di tornare all’IDE mantenendo il simulatore aperto.

#     Controlli finali

#     In Output è presente l’esito Build completed without errors.
#     Il target PC-Simulator è in stato verde.
#     La barra di stato dell’IDE riporta Connected to PC-Simulator, Synchronized.
#     Il browser mostra STone - Web simulator su 127.0.0.1 e visualizza:
#     barra laterale con le voci (Toggle…, Alarms, Screenshots, Credits);
#     pannelli I/O (DIGITAL INPUT, DIGITAL OUTPUT, LEDS, SEVEN SEGMENTS/UNIVERSAL);
#     una schermata dispositivo (es. Language: EN o Info - Circuit 1).
#     Dati da verificare
#     Nome esatto del menu/voce che avvia Full build (posizione e percorso nell’IDE).
#     Nome esatto del comando/pulsante “icona schermo” e la sua posizione nell’IDE.
#     Presenza e modalità di comando Stop specifico del simulatore/target (non indicato esplicitamente).
#     Procedura di salvataggio (progetto, configurazioni simulatore, screenshot) e relativi percorsi/file generati (non specificati).
#     Funzioni di debug/monitoraggio nell’IDE (watch, breakpoints, trace) e loro utilizzo con PC-Simulator (non descritte).
#     Condizioni che rendono attivo il pulsante SEND in “MULTI KEYS SEND” e formato/valori ammessi per Key code-1/Key code-2.
#     Troubleshooting dettagliato oltre al caso “X rossa su Connect to Target perché non compilato/downloadato” (ulteriori casi non forniti).
#     """
